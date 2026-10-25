"""Libros de entrada y salida. No cambia archivos de origen."""
from datetime import datetime, date
from io import BytesIO
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from .config import AVISO

def workbook_bytes(sheets, subtitle="Centro Demostrativo Huellas · Corte 2026-03-31"):
    """sheets: {nombre: (columnas, filas)}; cabeceras en fila 3."""
    wb = Workbook()
    wb.remove(wb.active)
    for index, (name, (columns, rows)) in enumerate(sheets.items()):
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.append([AVISO])
        ws.append([subtitle])
        ws.append(columns)
        ws["A1"].font = Font(name="Arial", size=12, bold=True, color="173F63")
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color="526479")
        ws.row_dimensions[1].height = 23
        ws.row_dimensions[2].height = 22
        for row in rows:
            ws.append(list(row))
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.font = Font(name="Arial", size=11, color="172D43")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                # Treat arbitrary input as literal text, never as executable Excel formulas.
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
                if isinstance(cell.value, (datetime, date)):
                    cell.number_format = "yyyy-mm-dd"
            ws.row_dimensions[row[0].row].height = 34 if name in ("Diccionario", "Hallazgos", "Reglas", "Conciliacion") else 24
        for cell in ws[3]:
            cell.fill = PatternFill("solid", fgColor="173F63")
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[3].height = 32
        for i, col in enumerate(columns, 1):
            width = {"explicacion": 65, "accion_sugerida": 62, "valores_permitidos": 62,
                     "puede_estar_vacio": 60, "observaciones": 44, "motivo": 70,
                     "combinacion": 44, "valor_observado": 28, "archivo": 29,
                     "criterio": 90, "nombre": 34, "gravedad": 30}.get(col, max(19, len(col) + 3))
            ws.column_dimensions[get_column_letter(i)].width = width
        if len(columns) < 4:
            ws.column_dimensions["A"].width = 45
            if len(columns) > 1:
                ws.column_dimensions["B"].width = 55
        # Height follows the actual wrapped text rather than clipping long explanations.
        for row in ws.iter_rows(min_row=4):
            lines=max(max(1,math.ceil(len(str(c.value or ""))/(ws.column_dimensions[c.column_letter].width*0.9))) for c in row)
            ws.row_dimensions[row[0].row].height=max(ws.row_dimensions[row[0].row].height,lines*15+8)
        ws.freeze_panes = "C4" if name == "Registros" else "A4"
        if rows:
            tab = Table(displayName=f"Tabla{index+1}", ref=f"A3:{get_column_letter(len(columns))}{len(rows)+3}")
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tab)
        ws.print_title_rows = "1:3"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    wb.properties.creator = "Vali-IA"
    wb.properties.title = "Vali-IA · Datos ficticios"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
