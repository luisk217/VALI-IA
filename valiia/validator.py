"""Reglas deterministas. Este módulo nunca lee el archivo de errores esperados."""
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from time import perf_counter
import hashlib
import re
import zipfile
import pandas as pd
from openpyxl import load_workbook
from .config import *

def vacio(value):
    return value is None or (isinstance(value, str) and not value.strip()) or (not isinstance(value, str) and pd.isna(value))

def visible(value):
    if vacio(value):
        return "(vacío)"
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)

def fecha(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for pattern, fmt in [(r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d"), (r"\d{2}/\d{2}/\d{4}", "%d/%m/%Y")]:
            if re.fullmatch(pattern, value):
                try:
                    return datetime.strptime(value, fmt).date()
                except ValueError:
                    return None
    return None

def entero(value, maximum=None):
    if vacio(value) or isinstance(value, (bool, date, datetime)):
        return None
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number < 0 or number != number.to_integral_value():
            return None
        if maximum is not None and number > maximum:
            return None
        return int(number)
    except (InvalidOperation, ValueError, TypeError):
        return None

def combo(key):
    return " / ".join(key)

def validar(data: bytes, archivo="archivo.xlsx"):
    start = perf_counter()
    result = {"aviso": AVISO, "version": "1.0.0", "archivo": archivo,
              "sha256": hashlib.sha256(data).hexdigest(), "periodo": "2026-01-01 a 2026-03-31",
              "corte": CORTE.isoformat(), "estructura_valida": False,
              "hallazgos": [], "conciliacion": [], "no_agrupables": [], "metricas": {}}
    findings = result["hallazgos"]

    def add(rule, code, severity, sheet, row, field, value, explanation, action, rid="", key=""):
        findings.append(dict(zip(HALLAZGOS, [f"H{len(findings)+1:03}", rule, code, severity,
            archivo, sheet, row, key, visible(rid) if not vacio(rid) else "", field,
            visible(value), explanation, action])))

    def structural(code, sheet, field, explanation, row=None):
        add("E00", code, "error", sheet, row, field, "", explanation,
            "Use una base de demostración y conserve sus hojas y cabeceras en la fila 3.")

    def finish(n=None):
        errors = {f["fila_excel"] for f in findings if f["hoja"] == "Registros" and f["regla"] != "E00" and f["gravedad"] == "error"}
        result["metricas"] = {
            "registros_revisados": n, "registros_con_error": len(errors) if n is not None else None,
            "porcentaje_sin_errores": round(100 * (n-len(errors))/n, 2) if n else None,
            "alertas_error": sum(f["gravedad"] == "error" for f in findings),
            "alertas_advertencia": sum(f["gravedad"] == "advertencia" for f in findings),
            "alertas_informacion": sum(f["gravedad"] == "información" for f in findings),
            "registros_no_agrupables": len(result["no_agrupables"]),
        }
        result["listo_para_informe"] = bool(result["estructura_valida"] and not any(f["gravedad"] in ("error", "advertencia") for f in findings))
        result["segundos_procesamiento"] = round(perf_counter()-start, 6)
        return result

    if not data:
        structural("archivo_vacio", "Archivo", "archivo", "El archivo está vacío. Seleccione un libro .xlsx con datos.")
        return finish()
    if len(data) > 5 * 1024 * 1024:
        structural("limite", "Archivo", "archivo", "La demostración admite archivos de hasta 5 MB.")
        return finish()
    try:
        with zipfile.ZipFile(BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist()) > 30 * 1024 * 1024:
                raise ValueError("El contenido descomprimido supera 30 MB")
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
        for sheet, cols in ESQUEMA.items():
            if sheet not in wb.sheetnames:
                structural("hoja_ausente", sheet, "hoja", f"Falta la hoja obligatoria {sheet}.")
                continue
            ws = wb[sheet]
            if ws.max_row > 10003 or ws.max_column > 50:
                structural("limite", sheet, "dimensiones", "Se admiten como máximo 10.000 filas de datos y 50 columnas por hoja.")
                continue
            headers = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
            for col in cols:
                if col not in headers:
                    structural("columna_ausente", sheet, col, f"Falta la columna {col} en la fila 3.", 3)
            nonempty = [h for h in headers if h is not None]
            if len(nonempty) != len(set(nonempty)):
                structural("cabecera_duplicada", sheet, "cabeceras", "Las cabeceras deben ser únicas.", 3)
            for cells in ws.iter_rows(min_row=4):
                for c in cells:
                    if c.data_type == "f":
                        structural("formula_no_admitida", sheet, c.coordinate, "La versión 1.0 requiere valores, no fórmulas ni sus resultados almacenados.", c.row)
        wb.close()
        if findings:
            return finish()
        frames = pd.read_excel(BytesIO(data), sheet_name=list(ESQUEMA), header=2,
                               dtype=object, keep_default_na=False, engine="openpyxl")
    except Exception as exc:
        structural("archivo_ilegible", "Archivo", "archivo", f"No se pudo leer el libro .xlsx ({type(exc).__name__}). Compruebe que no esté dañado ni protegido con contraseña.")
        return finish()

    cat = frames["Catalogos"]
    provided = {(visible(r.campo), visible(r.valor_permitido)) for r in cat.itertuples()}
    expected = {(k, v) for k, values in CATALOGOS.items() for v in values}
    if provided != expected or len(cat) != len(expected):
        structural("catalogo_invalido", "Catalogos", "campo / valor_permitido", "Los catálogos no corresponden a la versión 1.0. Deben contener exactamente las categorías documentadas, sin duplicados.")
    documented = {(str(r.hoja), str(r.campo)) for r in frames["Diccionario"].itertuples()}
    if not {(s,c) for s in ("Registros","Resumen_mensual") for c in ESQUEMA[s]}.issubset(documented):
        structural("diccionario_incompleto", "Diccionario", "hoja / campo", "El diccionario debe describir todos los campos de Registros y Resumen_mensual.")
    records = frames["Registros"]
    if records.empty:
        structural("sin_registros", "Registros", "filas", "La hoja Registros no contiene actividades. No se calcula un porcentaje de calidad.")
    if findings:
        return finish()
    result["estructura_valida"] = True
    # The uploaded catalog is used after checking its declared, closed version.
    catalogs = {k: set(cat.loc[cat.campo == k, "valor_permitido"]) for k in CATALOGOS}
    ids = Counter(visible(v) for v in records.id_registro if not vacio(v))
    counts = Counter()
    for i, r in enumerate(records.to_dict("records"), 4):
        rid = r["id_registro"]
        def reg(rule, code, field, reason, action="Verifique el dato en la fuente ficticia y proponga el ajuste para revisión humana.", severity="error"):
            add(rule, code, severity, "Registros", i, field, r[field], reason, action, rid)
        for col in OBLIGATORIOS:
            if vacio(r[col]):
                reg("R01", "obligatorio", col, f"El campo {col} es obligatorio y está vacío.")
        if not vacio(rid) and ids[visible(rid)] > 1:
            reg("R02", "duplicado", "id_registro", "El identificador aparece en varias filas; no se decide cuál conservar.")
        for col, allowed in catalogs.items():
            if not vacio(r[col]) and r[col] not in allowed:
                reg("R03", "categoria", col, "El valor no coincide exactamente con una categoría permitida.")
        activity = fecha(r["fecha_actividad"])
        in_period = activity is not None and INICIO <= activity <= CORTE
        if not vacio(r["fecha_actividad"]) and not in_period:
            reg("R04", "fecha_actividad", "fecha_actividad", "Fecha no interpretable o fuera del intervalo 2026-01-01 a 2026-03-31.")
        if not vacio(r["edad_meses"]) and entero(r["edad_meses"], 300) is None:
            reg("R05", "edad", "edad_meses", "La edad debe ser un número entero entre 0 y 300 meses.")
        follow = r["fecha_seguimiento"]
        if vacio(follow):
            if r["seguimiento_requerido"] == "Sí":
                reg("R06", "pendiente", "fecha_seguimiento", "Seguimiento pendiente al corte; no es un error ni implica incumplimiento de un plazo.", "Confirmar si el seguimiento se realizó y registrar su fecha cuando corresponda.", "información")
        else:
            fd = fecha(follow)
            if fd is None:
                reg("R06", "fecha_no_valida", "fecha_seguimiento", "La fecha de seguimiento no puede interpretarse.")
            else:
                if fd > CORTE:
                    reg("R06", "posterior_corte", "fecha_seguimiento", "El seguimiento informado es posterior al corte fijo.")
                if activity is not None and fd < activity:
                    reg("R06", "anterior_actividad", "fecha_seguimiento", "El seguimiento está fechado antes de la actividad.")
                if activity is None:
                    reg("R06", "orden_no_evaluable", "fecha_seguimiento", "No puede comprobarse el orden de las fechas porque falta una fecha de actividad válida.", severity="advertencia")
            if r["seguimiento_requerido"] == "No":
                reg("R06", "contradiccion", "fecha_seguimiento", "Existe una fecha de seguimiento aunque seguimiento_requerido es No.")
        if in_period and r["especie"] in catalogs["especie"] and r["tipo_actividad"] in catalogs["tipo_actividad"]:
            counts[(activity.strftime("%Y-%m"), r["especie"], r["tipo_actividad"])]+=1
        else:
            result["no_agrupables"].append({"fila_excel": i, "id_registro": visible(rid), "motivo": "Fecha, especie o tipo de actividad ausente, inválido o fuera del periodo."})
    record_errors = [f for f in findings if f["gravedad"] == "error"]
    unreliable = bool(record_errors)
    reasons = ", ".join(sorted({f["regla"] for f in record_errors}))
    summary = frames["Resumen_mensual"]
    groups = {}
    invalid_rows = set()
    for i, r in enumerate(summary.to_dict("records"), 4):
        key = tuple(visible(r[c]) for c in RESUMEN[:3])
        if key in COMBINACIONES:
            groups.setdefault(key, []).append((i, r))
        for field, allowed in [("periodo", PERIODOS), ("especie", CATALOGOS["especie"]), ("tipo_actividad", CATALOGOS["tipo_actividad"])]:
            if r[field] not in allowed:
                add("R08", "clave_invalida", "error", "Resumen_mensual", i, field, r[field], "Periodo o categoría inválida en el resumen.", "Restituir la combinación correcta tras revisar la fuente.", key=combo(key))
                invalid_rows.add(i)
        if entero(r["total_reportado"]) is None:
            add("R08", "total_invalido", "error", "Resumen_mensual", i, "total_reportado", r["total_reportado"], "El total debe ser un entero no negativo; un vacío no equivale a cero.", "Confirmar el total reportado y registrar un entero no negativo.", key=combo(key))
            invalid_rows.add(i)
    for key in COMBINACIONES:
        entries = groups.get(key, [])
        keytext = combo(key)
        total = None
        valid_summary = len(entries) == 1 and entries[0][0] not in invalid_rows
        if not entries:
            add("R08", "combinacion_ausente", "error", "Resumen_mensual", None, "combinacion", keytext, "Falta una combinación esperada, incluso si su total es cero.", "Añadir la combinación y verificar su total.", key=keytext)
        elif len(entries) > 1:
            for row, r in entries:
                add("R08", "combinacion_duplicada", "error", "Resumen_mensual", row, "combinacion", keytext, "La combinación del resumen aparece varias veces.", "Revisar las filas; no sumarlas ni escoger una automáticamente.", key=keytext)
        if valid_summary:
            row, r = entries[0]
            total = entero(r["total_reportado"])
            difference = total-counts[key]
            state = ("Provisional: " if unreliable else "") + ("diferencia" if difference else "coincide")
            if difference:
                add("R07", "diferencia", "advertencia" if unreliable else "error", "Resumen_mensual", row, "total_reportado", total,
                    f"Reportado {total}, recuento de filas agrupables {counts[key]}, diferencia {difference:+d}." + (" Comparación provisional por errores en Registros." if unreliable else ""),
                    "Revisar primero los registros y luego conciliar el total; no ajustar automáticamente.", key=keytext)
        else:
            difference, state = None, "No comparable: resumen inválido"
        result["conciliacion"].append({"periodo": key[0], "especie": key[1], "tipo_actividad": key[2],
            "total_reportado": total, "recuento_filas_agrupables": counts[key], "diferencia": difference,
            "estado": state, "motivo": (f"Errores en Registros: {reasons}. Incluye duplicados y filas agrupables con otros errores; {len(result['no_agrupables'])} filas no asignables se detallan aparte." if unreliable else "Recuento de actividades, no de animales únicos.")})
    return finish(len(records))
