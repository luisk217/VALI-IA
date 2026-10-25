import hashlib
from io import BytesIO
import json
from pathlib import Path
import unittest
import zipfile
from openpyxl import load_workbook
from valiia.validator import validar
from valiia.reporting import export_excel, ai_package
from generar_datos import build_data

ROOT=Path(__file__).resolve().parents[1]
FIELDS=["regla","codigo","gravedad","hoja","fila_excel","campo","combinacion"]
def key(item):
    return tuple(item.get(k) for k in FIELDS)
def base(name="base_correcta.xlsx"):
    return (ROOT/"datos"/name).read_bytes()
def mutated(callback):
    wb=load_workbook(BytesIO(base()))
    callback(wb)
    out=BytesIO()
    wb.save(out)
    return out.getvalue()

class ValidatorTests(unittest.TestCase):
    def test_correcta(self):
        r=validar(base())
        self.assertTrue(r["listo_para_informe"])
        self.assertEqual(r["hallazgos"],[])
        self.assertEqual(r["metricas"]["registros_revisados"],200)
        self.assertEqual(sum(c["recuento_filas_agrupables"] for c in r["conciliacion"]),200)
        self.assertEqual(len(r["conciliacion"]),18)

    def test_oraculo_independiente(self):
        expected=json.loads((ROOT/"datos"/"errores_esperados.json").read_text(encoding="utf-8"))
        r=validar(base("base_con_errores.xlsx"))
        obtained=[f for f in r["hallazgos"] if f["gravedad"]!="información"]
        self.assertCountEqual([key(f) for f in obtained],[key(f) for f in expected["hallazgos_esperados"]])
        self.assertEqual(r["metricas"]["alertas_informacion"],1)
        self.assertEqual(r["metricas"]["registros_con_error"],18)
        self.assertEqual(r["metricas"]["alertas_error"],24)
        self.assertEqual(r["metricas"]["porcentaje_sin_errores"],91.0)
        self.assertEqual(len(r["no_agrupables"]),4)
        self.assertEqual(sum(c["recuento_filas_agrupables"] for c in r["conciliacion"])+len(r["no_agrupables"]),200)
        self.assertTrue(all(c["estado"].startswith(("Provisional","No comparable")) for c in r["conciliacion"]))

    def test_excepciones_validas(self):
        r=validar(base("base_excepciones_validas.xlsx"))
        self.assertEqual(r["metricas"]["alertas_error"],0)
        self.assertEqual(r["metricas"]["alertas_advertencia"],0)
        self.assertEqual(r["metricas"]["alertas_informacion"],3)
        self.assertEqual(r["metricas"]["porcentaje_sin_errores"],100)

    def test_columna_ausente(self):
        r=validar(mutated(lambda wb: wb["Registros"].delete_cols(6)))
        self.assertFalse(r["estructura_valida"])
        self.assertTrue(any(f["codigo"]=="columna_ausente" and f["campo"]=="zona" for f in r["hallazgos"]))
        self.assertIsNone(r["metricas"]["porcentaje_sin_errores"])

    def test_hoja_ausente(self):
        r=validar(mutated(lambda wb: wb.remove(wb["Resumen_mensual"])))
        self.assertEqual(r["hallazgos"][0]["codigo"],"hoja_ausente")

    def test_archivo_vacio(self):
        r=validar(b"")
        self.assertEqual(r["hallazgos"][0]["codigo"],"archivo_vacio")
        self.assertIsNone(r["metricas"]["registros_revisados"])

    def test_archivo_ilegible(self):
        r=validar(b"esto no es un Excel")
        self.assertEqual(r["hallazgos"][0]["codigo"],"archivo_ilegible")

    def test_hoja_sin_datos(self):
        r=validar(mutated(lambda wb: wb["Registros"].delete_rows(4,200)))
        self.assertFalse(r["estructura_valida"])
        self.assertIsNone(r["metricas"]["porcentaje_sin_errores"])

    def test_varias_actividades_mismo_animal(self):
        wb=load_workbook(BytesIO(base()))
        values=[wb["Registros"].cell(i,2).value for i in range(4,204)]
        self.assertLess(len(set(values)),200)
        self.assertFalse(any(f["regla"]=="R02" for f in validar(base())["hallazgos"]))

    def test_total_diferente_sin_errores_registros(self):
        def change(wb):
            wb["Resumen_mensual"]["D4"].value+=7
        r=validar(mutated(change))
        self.assertEqual(len(r["hallazgos"]),1)
        self.assertEqual(r["hallazgos"][0]["regla"],"R07")
        self.assertEqual(r["hallazgos"][0]["gravedad"],"error")
        self.assertEqual(r["conciliacion"][0]["diferencia"],7)
        self.assertEqual(r["metricas"]["registros_con_error"],0)

    def test_cero_no_es_vacio(self):
        r=validar(base())
        zero=[c for c in r["conciliacion"] if c["total_reportado"]==0]
        self.assertTrue(zero)
        def change(wb):
            wb["Resumen_mensual"]["D21"]=None
        empty=validar(mutated(change))
        self.assertTrue(any(f["codigo"]=="total_invalido" for f in empty["hallazgos"]))
        self.assertIsNone(empty["conciliacion"][-1]["total_reportado"])

    def test_no_modifica_origen(self):
        path=ROOT/"datos"/"base_con_errores.xlsx"
        original=path.read_bytes()
        result=validar(original)
        self.assertEqual(path.read_bytes(),original)
        self.assertEqual(result["sha256"],hashlib.sha256(original).hexdigest())

    def test_exportaciones(self):
        r=validar(base("base_con_errores.xlsx"))
        wb=load_workbook(BytesIO(export_excel(r)))
        self.assertEqual(wb["Hallazgos"].max_row,len(r["hallazgos"])+3)
        with zipfile.ZipFile(BytesIO(ai_package(r))) as z:
            decoded=json.loads(z.read("resultados_validacion.json"))
            self.assertEqual(decoded["metricas"],r["metricas"])
            self.assertIn("manual",z.read("LEEME.txt").decode().lower())

    def test_semilla_reproducible(self):
        first,oracle1=build_data()
        second,oracle2=build_data()
        self.assertEqual(first,second)
        self.assertEqual(oracle1,oracle2)

    def test_catalogo_alterado(self):
        r=validar(mutated(lambda wb: setattr(wb["Catalogos"]["B4"],"value","Ave")))
        self.assertFalse(r["estructura_valida"])
        self.assertTrue(any(f["codigo"]=="catalogo_invalido" for f in r["hallazgos"]))

    def test_formula_no_se_ejecuta(self):
        r=validar(mutated(lambda wb: setattr(wb["Registros"]["H4"],"value","=10+10")))
        self.assertFalse(r["estructura_valida"])
        self.assertTrue(any(f["codigo"]=="formula_no_admitida" for f in r["hallazgos"]))

    def test_ausencia_no_genera_errores_en_cascada(self):
        r=validar(mutated(lambda wb: setattr(wb["Registros"]["H4"],"value",None)))
        matching=[f for f in r["hallazgos"] if f["campo"]=="edad_meses"]
        self.assertEqual([f["regla"] for f in matching],["R01"])

    def test_fechas_limites_y_orden_no_evaluable(self):
        def change(wb):
            wb["Registros"]["C4"]="31/03/2026"
            wb["Registros"]["J4"]="31/03/2026"
        r=validar(mutated(change))
        self.assertFalse(any(f["regla"] in ["R04","R06"] for f in r["hallazgos"]))
        r=validar(mutated(lambda wb: setattr(wb["Registros"]["C4"],"value","mal")))
        self.assertTrue(any(f["codigo"]=="orden_no_evaluable" for f in r["hallazgos"]))

class InterfaceTests(unittest.TestCase):
    def test_flujo_interfaz(self):
        from streamlit.testing.v1 import AppTest
        at=AppTest.from_file(str(ROOT/"app.py"),default_timeout=30).run()
        self.assertEqual(len(at.exception),0)
        at.sidebar.button[0].click().run()
        self.assertEqual(len(at.exception),0)
        self.assertEqual(at.metric[0].value,"200")
        self.assertEqual(at.metric[1].value,"18")
        # Filtering must not erase results or change global metrics.
        at.multiselect[0].set_value(["información"]).run()
        self.assertEqual(at.metric[1].value,"18")
        self.assertEqual(len(at.dataframe[0].value),1)
        # Switching a source invalidates stale results before pressing Validate.
        at.sidebar.selectbox[0].select("Base correcta").run()
        self.assertEqual(len(at.metric),0)
        at.sidebar.button[0].click().run()
        self.assertEqual(at.metric[1].value,"0")
        self.assertEqual(len(at.exception),0)
        at.sidebar.selectbox[0].select("Excepciones válidas").run()
        at.sidebar.button[0].click().run()
        self.assertEqual(at.metric[2].value,"100.0%")
        self.assertEqual(len(at.exception),0)

if __name__=="__main__":
    unittest.main(verbosity=2)
